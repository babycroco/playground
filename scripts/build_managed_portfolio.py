"""
scripts/build_managed_portfolio.py — Parse Managers sheet → managed_portfolio.csv

Reads data/Rental_Payments-NEW.xlsx (Managers sheet), normalises street and
unit identifiers via audit.unit_matcher, and writes data/managed_portfolio.csv.

Output CSV columns:
    owner, raw_property_name, street, unit_number, canonical_key,
    hg_2024, hg_2025, hg_2026, latest_hg, latest_hg_year,
    is_direct_payer, needs_wp_update

Run with:
    python3 -m scripts.build_managed_portfolio
"""

import os
import re
import csv

import openpyxl

from audit.unit_matcher import normalize_street, normalize_unit, make_canonical_key

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
INPUT_PATH = os.path.join(DATA_DIR, "Rental_Payments-NEW.xlsx")
OUTPUT_PATH = os.path.join(DATA_DIR, "managed_portfolio.csv")

OUTPUT_COLUMNS = [
    "owner",
    "raw_property_name",
    "street",
    "unit_number",
    "canonical_key",
    "hg_2024",
    "hg_2025",
    "hg_2026",
    "latest_hg",
    "latest_hg_year",
    "is_direct_payer",
    "needs_wp_update",
]

# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

_SAFE_FORMULA_RE = re.compile(r"^[\d\s\+\-\*\/\.\(\)]+$")


def eval_formula(cell_value) -> "float | None":
    """
    If cell_value is a string starting with '=', parse the rest as arithmetic.

    Only digits, spaces, +, -, *, /, ., () are allowed — no other chars.
    Uses Python eval() after validation.
    Returns None if the expression is empty, unsafe, or raises an error.
    """
    if not isinstance(cell_value, str):
        return None
    s = cell_value.strip()
    if not s.startswith("="):
        return None
    expr = s[1:].strip()
    if not expr:
        return None
    if not _SAFE_FORMULA_RE.match(expr):
        return None
    try:
        result = eval(expr, {"__builtins__": {}})  # noqa: S307
        return float(result)
    except Exception:
        return None


def parse_hg_value(cell_value) -> "tuple[float | None, bool]":
    """
    Returns (value, is_direct_payer).

    - Numeric               → (float(value), False)
    - Formula string "=…"   → (eval_formula(value), False)
    - "pay on their own" /
      "eigentümer zahlt"    → (None, True)
    - "LEER"                → (None, False)   [vacant — still needs audit]
    - None / empty          → (None, False)
    """
    if cell_value is None:
        return (None, False)

    # Numeric (int or float coming from openpyxl)
    if isinstance(cell_value, (int, float)):
        return (float(cell_value), False)

    if not isinstance(cell_value, str):
        # Try coercing other types
        try:
            return (float(cell_value), False)
        except (TypeError, ValueError):
            return (None, False)

    s = cell_value.strip()

    if not s:
        return (None, False)

    # Formula
    if s.startswith("="):
        return (eval_formula(s), False)

    # Direct payer markers (case-insensitive)
    lower = s.lower()
    if "pay on their own" in lower or "eigentümer zahlt" in lower:
        return (None, True)

    # Vacant marker
    if s.upper() == "LEER":
        return (None, False)

    # Attempt numeric coercion for any remaining strings (e.g. "320.0")
    try:
        return (float(s), False)
    except ValueError:
        return (None, False)


def parse_property_name(raw: str) -> "tuple[str, str]":
    """
    Parse a raw property name cell into (street_raw, unit_raw).

    Steps:
      1. Strip everything after the first " / " or "/" — multi-address cells;
         take only the first address segment.
      2. Find WE number with regex WE\\s*(\\d+) (case-insensitive).
      3. Street = everything before the WE match, stripped of trailing
         whitespace and dots.

    Returns (street_raw, unit_raw) e.g. ("Residenzstr. 129", "15").
    Both may be empty strings if the input is empty or unparseable.
    """
    if not raw or not isinstance(raw, str):
        return ("", "")

    # 1. Take only the first address segment (strip after " / " or "/")
    #    Use " / " first so that "Weserstr. 59 WE19/ Wildenbruch …" is split
    #    at the "/" after the unit number.
    segment = re.split(r"\s*/\s*", raw, maxsplit=1)[0].strip()

    # 2. Find WE number
    m = re.search(r"WE\s*(\d+)", segment, flags=re.IGNORECASE)
    if not m:
        # No unit found — return the whole segment as street, unit empty
        return (segment.rstrip(" ."), "")

    unit_raw = m.group(1)

    # 3. Street = everything before the WE match
    street_raw = segment[: m.start()].rstrip(" .")

    return (street_raw, unit_raw)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_managed_portfolio():
    wb = openpyxl.load_workbook(INPUT_PATH, data_only=False)

    if "Managers" not in wb.sheetnames:
        raise ValueError(f"'Managers' sheet not found in {INPUT_PATH}")

    ws = wb["Managers"]

    seen_keys: set[str] = set()
    portfolio: list[dict] = []

    # Counters for summary
    count_with_2026 = 0
    count_missing_2026 = 0
    count_direct_payer = 0
    count_leer = 0

    rows = list(ws.iter_rows(values_only=True))
    # Skip header row (row 0)
    for row in rows[1:]:
        # Unpack columns A–F (indices 0–5); pad if row is shorter
        owner = row[0] if len(row) > 0 else None
        raw_property = row[1] if len(row) > 1 else None
        raw_2024 = row[2] if len(row) > 2 else None
        raw_2025 = row[3] if len(row) > 3 else None
        raw_2026 = row[4] if len(row) > 4 else None
        raw_2027 = row[5] if len(row) > 5 else None  # noqa: F841 (reserved for future use)

        # Skip rows where Property name is empty
        if not raw_property or (isinstance(raw_property, str) and not raw_property.strip()):
            continue

        # Parse property name
        street_raw, unit_raw = parse_property_name(str(raw_property))

        # Normalize via audit.unit_matcher
        street = normalize_street(street_raw) if street_raw else ""
        unit_number = normalize_unit(unit_raw) if unit_raw else ""
        canonical_key = make_canonical_key(street_raw, unit_raw)

        # Parse HG values
        hg_2024, dp_2024 = parse_hg_value(raw_2024)
        hg_2025, dp_2025 = parse_hg_value(raw_2025)
        hg_2026, dp_2026 = parse_hg_value(raw_2026)

        is_direct_payer = dp_2024 or dp_2025 or dp_2026

        # Check for LEER specifically in 2026
        is_leer = (
            isinstance(raw_2026, str)
            and raw_2026.strip().upper() == "LEER"
        )

        # Determine latest non-None HG: prefer 2026 > 2025 > 2024
        latest_hg = None
        latest_hg_year = None
        for value, year in [(hg_2026, 2026), (hg_2025, 2025), (hg_2024, 2024)]:
            if value is not None:
                latest_hg = value
                latest_hg_year = year
                break

        needs_wp_update = hg_2026 is None

        # Deduplicate by canonical_key — keep first occurrence
        if canonical_key and canonical_key in seen_keys:
            continue
        if canonical_key:
            seen_keys.add(canonical_key)

        # Update counters
        if is_direct_payer:
            count_direct_payer += 1
        elif is_leer:
            count_leer += 1
        elif hg_2026 is not None:
            count_with_2026 += 1
        else:
            count_missing_2026 += 1

        portfolio.append({
            "owner": owner or "",
            "raw_property_name": raw_property,
            "street": street_raw,          # human-readable, e.g. "Residenzstr. 129"
            "unit_number": unit_number,
            "canonical_key": canonical_key,
            "hg_2024": hg_2024 if hg_2024 is not None else "",
            "hg_2025": hg_2025 if hg_2025 is not None else "",
            "hg_2026": hg_2026 if hg_2026 is not None else "",
            "latest_hg": latest_hg if latest_hg is not None else "",
            "latest_hg_year": latest_hg_year if latest_hg_year is not None else "",
            "is_direct_payer": is_direct_payer,
            "needs_wp_update": needs_wp_update,
        })

    # Write CSV
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(OUTPUT_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(portfolio)

    total = len(portfolio)
    print(f"Managed portfolio: {total} units")
    print(f"  - With 2026 HG: {count_with_2026} units")
    print(f"  - Missing 2026 HG: {count_missing_2026} units (need WP extraction)")
    print(f"  - Direct payers: {count_direct_payer} units (excluded from audit)")
    print(f"  - LEER (vacant): {count_leer} units")
    print(f"Output: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_managed_portfolio()
